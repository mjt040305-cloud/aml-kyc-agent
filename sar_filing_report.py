"""
sar_filing_report.py
---------------------
Generates the final SAR filing document (Word/.docx) for a fully
completed (two-person signed-off) case, formatted for submission to the
Financial Intelligence Unit (FIU) Zimbabwe - the statutory recipient of
Suspicious Transaction Reports under the Money Laundering and Proceeds of
Crime Act.

This module only produces output for transactions that have BOTH the
first officer's decision AND a second officer's co-signature recorded
(see case_store.py's two-person sign-off workflow). A case still awaiting
co-signature has no completed filing package to produce - see
build_sar_filing_docx()'s early return in that situation.

Word (.docx) rather than PDF, so the receiving institution/FIU can
annotate, copy figures from, or merge this into their own filing
template - this uses python-docx, the Python library for generating
.docx files programmatically (not to be confused with the separate
npm "docx" package used for document authoring in other contexts).
"""

from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

import sar_narrative

NAVY = RGBColor(0x1F, 0x38, 0x64)


def build_sar_filing_docx(case: dict, output_path: str):
    """
    `case`: the dict returned by case_store.load_case() for a case whose
    status is 'completed'.

    Returns (output_path, error) - error is None on success. If the case
    has no fully co-signed SAR escalations, returns (None, message)
    instead of producing an empty or misleading filing document.
    """
    cosign = case.get("cosign") or {}
    escalated = {
        tid: e for tid, e in cosign.items()
        if e.get("decision") == "Escalate to SAR filing" and e.get("second_officer_id")
    }
    if not escalated:
        return None, "This case has no fully co-signed SAR escalations to file."

    final_report = case.get("final_report") or []
    txn_by_id = {t["transaction_id"]: t for t in final_report}

    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Suspicious Activity Report \u2014 Filing Package")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run("Prepared for submission to the Financial Intelligence Unit (FIU), Zimbabwe")
    srun.italic = True
    srun.font.size = Pt(10)
    srun.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run(f"Case Reference: {case['case_id']}").bold = True
    doc.add_paragraph(f"Prepared: {datetime.now().strftime('%Y-%m-%d %H:%M')}    |    Transactions filed: {len(escalated)}")
    doc.add_paragraph()

    for tid, entry in escalated.items():
        txn = txn_by_id.get(tid, {})

        heading = doc.add_heading(level=2)
        heading.add_run(f"Transaction {tid}").font.color.rgb = NAVY

        oc = txn.get("original_currency", "USD")
        oa = txn.get("original_amount", txn.get("amount", 0)) or 0
        usd = txn.get("usd_equivalent", txn.get("amount", 0)) or 0

        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        fields = [
            ("Customer ID", txn.get("customer_id", "N/A")),
            ("Date", txn.get("date", "N/A")),
            ("Original amount", f"{oc} {oa:,.2f}"),
            ("USD equivalent", f"${usd:,.2f}"),
            ("FX rate source", txn.get("fx_rate_source", "N/A")),
            ("Counterparty jurisdiction", txn.get("counterparty_country", "N/A")),
            ("Overall risk score", f"{txn.get('risk_score', 'N/A')} ({txn.get('risk_bucket', 'N/A')})"),
        ]
        for label, value in fields:
            row = table.add_row()
            row.cells[0].text = label
            row.cells[0].paragraphs[0].runs[0].bold = True
            row.cells[1].text = str(value)

        triggered = txn.get("triggered_rules", [])
        if triggered:
            doc.add_paragraph()
            doc.add_paragraph("AML rules triggered:").runs[0].bold = True
            for r in triggered:
                doc.add_paragraph(f"{r.get('label', '')}: {r.get('reason', '')}", style="List Bullet")

        doc.add_paragraph()
        doc.add_paragraph("Narrative:").runs[0].bold = True
        narrative, _ = sar_narrative.generate_sar_narrative(txn)
        for para_text in (narrative or "").split("\n\n"):
            if para_text.strip():
                doc.add_paragraph(para_text.strip())

        doc.add_paragraph()
        doc.add_paragraph("Sign-off (two-person maker-checker):").runs[0].bold = True
        doc.add_paragraph(
            f"First reviewer: {entry.get('first_officer_name', 'N/A')} "
            f"({entry.get('first_officer_role', 'N/A')}) \u2014 {entry.get('first_signed_at', 'N/A')}"
        )
        doc.add_paragraph(
            f"Second reviewer (co-sign): {entry.get('second_officer_name', 'N/A')} "
            f"({entry.get('second_officer_role', 'N/A')}) \u2014 {entry.get('second_signed_at', 'N/A')}"
        )
        doc.add_paragraph()
        sep = doc.add_paragraph()
        sep.paragraph_format.space_after = Pt(12)

    doc.save(output_path)
    return output_path, None
