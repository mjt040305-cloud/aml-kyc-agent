"""
escalation_report.py
----------------------
Generates a Word (.docx) handoff document listing every transaction an
officer marked "Escalate to SAR filing" in a completed review - meant to
be sent to whoever is responsible for further investigation of those
transactions internally.

This is deliberately different from sar_filing_report.py:
- sar_filing_report.py only ever includes a transaction once it has BOTH
  the first officer's decision AND a second officer's co-signature (the
  final, FIU-ready filing package).
- escalation_report.py includes every transaction marked for escalation
  in THIS batch's review, regardless of whether two-person sign-off has
  completed yet - it's an internal "here's what needs further
  investigation" handoff, not the final regulator submission.
"""

from datetime import datetime
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

NAVY = RGBColor(0x1F, 0x38, 0x64)


def build_escalation_report_docx(final_report: list, case_id: str, output_path: str):
    """
    `final_report`: the list of transaction dicts from st.session_state.final_report
    (each already carrying review_status/reviewed_by/reviewer_notes from
    the human review step).

    Returns (output_path, error) - error is None on success. If nothing in
    this batch was marked for escalation, returns (None, message) instead
    of producing an empty document.
    """
    escalated = [t for t in final_report if t.get("review_status") == "Escalate to SAR filing"]
    if not escalated:
        return None, "No transactions in this batch were marked 'Escalate to SAR filing'."

    doc = Document()
    for section in doc.sections:
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run("Transactions Escalated for Further Investigation")
    run.bold = True
    run.font.size = Pt(18)
    run.font.color.rgb = NAVY

    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.CENTER
    srun = subtitle.add_run(
        "This agent flags transactions and provides evidence for review - it does not determine "
        "that any transaction is illegal. Each item below was escalated by the reviewing compliance "
        "officer for further investigation."
    )
    srun.italic = True
    srun.font.size = Pt(9.5)
    srun.font.color.rgb = RGBColor(0x60, 0x60, 0x60)

    doc.add_paragraph()
    meta = doc.add_paragraph()
    meta.add_run(f"Case Reference: {case_id}").bold = True
    doc.add_paragraph(f"Prepared: {datetime.now().strftime('%Y-%m-%d %H:%M')}    |    Transactions escalated: {len(escalated)}")
    doc.add_paragraph()

    for txn in escalated:
        heading = doc.add_heading(level=2)
        heading.add_run(f"Transaction {txn.get('transaction_id', 'N/A')}").font.color.rgb = NAVY

        oc = txn.get("original_currency", "USD")
        oa = txn.get("original_amount", txn.get("amount", 0)) or 0
        usd = txn.get("usd_equivalent", txn.get("amount", 0)) or 0

        table = doc.add_table(rows=0, cols=2)
        table.style = "Light Grid Accent 1"
        fields = [
            ("Customer ID", txn.get("customer_id", "N/A")),
            ("Date", txn.get("date", "N/A")),
            ("Original amount", f"{oc} {oa:,.2f}"),
            ("USD equivalent", f"${usd:,.2f}"),
            ("FX rate source", txn.get("fx_rate_source", "N/A")),
            ("Counterparty jurisdiction", txn.get("counterparty_country", "N/A")),
            ("Overall risk score", f"{txn.get('risk_score', 'N/A')} ({txn.get('risk_bucket', 'N/A')})"),
            ("Reviewed by", txn.get("reviewed_by", "N/A")),
            ("Reviewer notes", txn.get("reviewer_notes") or "(none)"),
        ]
        for label, value in fields:
            row = table.add_row()
            row.cells[0].text = label
            row.cells[0].paragraphs[0].runs[0].bold = True
            row.cells[1].text = str(value)

        triggered = txn.get("triggered_rules", [])
        if triggered:
            doc.add_paragraph()
            doc.add_paragraph("AML rules triggered:").runs[0].bold = True
            for r in triggered:
                doc.add_paragraph(f"{r.get('label', '')}: {r.get('reason', '')}", style="List Bullet")

        doc.add_paragraph()

    doc.save(output_path)
    return output_path, None
